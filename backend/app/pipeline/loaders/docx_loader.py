"""Word (docx) 文档加载器（基于 python-docx）

支持提取文本（含表格内文字）和嵌入图片。图片写入临时目录（避免内存压力），
由 pipeline 的 OCR 流程处理后清理。
"""

import hashlib
import io
import os
import tempfile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table

from app.pipeline.loader import BaseLoader, EmbeddedImage, LoadResult
from app.pipeline.loaders.docx_repair import repair_docx_bytes
from app.pipeline.loaders.docx_xml_text import paragraph_text, paragraphs_from_bytes

_W_P = qn("w:p")
_W_TBL = qn("w:tbl")


# 最小图片数据大小（字节），过小的图片跳过
_MIN_IMAGE_BYTES = 1024
# 单文档最大提取图片数量
_MAX_IMAGES_PER_DOC = 50


class DocxLoader(BaseLoader):
    """处理 .docx 文件的加载器，同时提取文本和嵌入图片"""

    def load(self, file_path: str) -> LoadResult:
        """加载 Word 文件，提取全部段落文本和嵌入图片

        图片写入临时目录，通过 content_hash 去重。

        Args:
            file_path: 文件路径

        Returns:
            LoadResult: 包含文件内容、元数据和嵌入图片列表

        Raises:
            FileNotFoundError: 文件不存在
            ValueError: 无效的 docx 文件
        """
        if not os.path.isfile(file_path):
            raise FileNotFoundError(f"文件不存在: {file_path}")

        file_size = os.path.getsize(file_path)

        doc, repairs = self._open_document(file_path)

        # 按文档顺序提取正文段落与表格文字
        blocks = self._iter_body_blocks(doc)
        if not "".join(blocks).strip():
            # 结构化遍历一无所获：包结构不典型（python-docx 连 body 都读不出东西）。
            # 退回直接解析 document.xml——宁可粗一点，也不能静默返回空正文。
            with open(file_path, "rb") as handle:
                blocks = paragraphs_from_bytes(handle.read())
        content = "\n\n".join(blocks)

        # 提取嵌入图片
        images = self._extract_images(doc)

        metadata = {
            "filename": os.path.basename(file_path),
            "file_type": "docx",
            "file_size": file_size,
            "embedded_image_count": len(images),
        }
        if repairs:
            # 自愈说明：这份文档的包声明被修过。供排障与入库核对使用。
            metadata["docx_repairs"] = repairs

        return LoadResult(content=content, metadata=metadata, images=images)

    @staticmethod
    def _open_document(file_path: str) -> tuple[Document, list[str]]:
        """打开 docx；包声明不自洽时先自愈再试一次。

        实测语料里有 51 份 python-docx 直接打不开的文档，正文却完全完好：44 份声明了
        不存在的 ``userCustomization/customUI.xml``，7 份是 macro-enabled（``.docm``
        改名）。这些是包的元数据问题，自愈（改内存副本、不动原文件）比判文档失败划算。

        Returns:
            ``(Document, 修复说明列表)``；首次即成功时说明列表为空。
        """
        try:
            return Document(file_path), []
        except Exception as first_error:
            try:
                with open(file_path, "rb") as handle:
                    repaired, notes = repair_docx_bytes(handle.read())
            except Exception as repair_error:  # noqa: BLE001 — 修不了就报原始错误
                raise ValueError(
                    f"无法解析 docx 文件: {file_path}，错误: {first_error}"
                    f"（自愈失败: {repair_error}）"
                ) from first_error
            if not notes:
                raise ValueError(
                    f"无法解析 docx 文件: {file_path}，错误: {first_error}"
                ) from first_error
            try:
                return Document(io.BytesIO(repaired)), notes
            except Exception as second_error:
                raise ValueError(
                    f"无法解析 docx 文件: {file_path}，错误: {first_error}；"
                    f"自愈后仍失败: {second_error}"
                ) from second_error

    @staticmethod
    def _iter_body_blocks(doc: Document) -> list[str]:
        """按 body 子元素顺序产出段落文本与表格文字。

        表格文字此前被整体丢弃——``doc.paragraphs`` 不含表格内段落。实测语料里有 37 份
        文档**整部法条都放在一张表格里**（最多的 403 段、约 1 万字），丢弃表格等于整篇
        丢失。这里按顺序遍历，段落与表格保持原有先后关系，它们落进同一条文本流，
        切分器照常按「第X条」切分。
        """
        return _blocks_of(doc.element.body, doc)

    @staticmethod
    def _extract_images(doc: Document) -> list[EmbeddedImage]:
        """从 Word 文档中提取所有嵌入图片，写入临时目录

        通过 content_hash 去重，过滤过小的图片。

        Args:
            doc: python-docx Document 对象

        Returns:
            提取到的 EmbeddedImage 列表
        """
        tmp_dir = tempfile.mkdtemp(prefix="docx_images_")
        images: list[EmbeddedImage] = []
        seen_hashes: set[str] = set()
        img_index = 0

        for rel in doc.part.rels.values():
            if "image" not in rel.reltype:
                continue

            if len(images) >= _MAX_IMAGES_PER_DOC:
                break

            try:
                image_part = rel.target_part
                image_bytes = image_part.blob
            except Exception:
                continue

            if not image_bytes or len(image_bytes) < _MIN_IMAGE_BYTES:
                continue

            # 计算 hash 去重
            img_hash = hashlib.md5(image_bytes).hexdigest()
            if img_hash in seen_hashes:
                continue
            seen_hashes.add(img_hash)

            # 从 content_type 推断格式
            content_type = getattr(image_part, "content_type", "")
            img_format = "png"
            if "jpeg" in content_type or "jpg" in content_type:
                img_format = "jpeg"
            elif "png" in content_type:
                img_format = "png"
            elif "gif" in content_type:
                img_format = "gif"
            elif "bmp" in content_type:
                img_format = "bmp"

            img_index += 1
            img_filename = f"docx_img{img_index}_{img_hash[:8]}.{img_format}"
            img_path = os.path.join(tmp_dir, img_filename)

            with open(img_path, "wb") as f:
                f.write(image_bytes)

            images.append(
                EmbeddedImage(
                    file_path=img_path,
                    format=img_format,
                    page_or_index=img_index,
                    content_hash=img_hash,
                    description=f"docx_img{img_index}",
                )
            )

        return images


def _table_texts(table: Table) -> list[str]:
    """取一张表格的文字，按行优先顺序、逐段落一项。

    合并单元格会被同一行的多个位置返回同一个 ``tc`` 元素，需要去重避免文字重复。

    **去重不能用裸 ``id(tc)``**：``cell._tc`` 是 lxml 按需创建的代理对象，代理被回收后
    新代理可以复用到同一个 ``id``，于是不同行的单元格会被误判成"同一个合并单元格"而
    整行丢弃（实测《吉林省个体工商户条例》125 行里只留下第 1 行）。这里把代理对象
    存进 ``alive`` 保活，``id`` 才可靠。
    """
    texts: list[str] = []
    seen: set[int] = set()
    alive: list = []
    for row in table.rows:
        for cell in row.cells:
            marker = id(cell._tc)
            if marker in seen:
                continue
            seen.add(marker)
            alive.append(cell._tc)
            texts.extend(_blocks_of(cell._tc, cell))
    return texts


def _blocks_of(parent_element, parent) -> list[str]:
    """按顺序取容器（body 或单元格）下的段落文本与表格文字。

    单元格里可能再套表格（实测 11 份文档是"表格套表格"，外层单元格本身没有文字），
    因此这里递归处理，而不是用 ``cell.text``——后者只看直接段落。
    """
    blocks: list[str] = []
    for child in parent_element.iterchildren():
        if child.tag == _W_P:
            # 用 docx_xml_text.paragraph_text 而不是 Paragraph.text：后者看不到
            # ``w:ins``（修订插入）与文本框里的 run，会让「第X条」的行首锚点失效。
            # 语义定义与兜底路径共用一处，见该模块的说明。
            text = paragraph_text(child).strip()
            if text:
                blocks.append(text)
        elif child.tag == _W_TBL:
            blocks.extend(_table_texts(Table(child, parent)))
    return blocks
