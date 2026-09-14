"""docx 加载器的容错：包自愈（``docx_repair``）与表格文字纳入正文。

覆盖实测语料里真实存在的三类问题：悬空关联声明（44 份）、macro-enabled 内容类型
（7 份）、正文整篇放在表格里（37 份）。fixture 用 python-docx 现造正常 docx，再按
真实语料的形态去污染它。
"""

from __future__ import annotations

import os as _os

_os.environ.setdefault("JWT_SECRET", "test-only-jwt-secret-not-for-production")

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document

from app.pipeline.loaders.docx_loader import DocxLoader
from app.pipeline.loaders.docx_repair import repair_docx_bytes

_WML_MAIN = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
_MACRO_MAIN = "application/vnd.ms-word.document.macroEnabled.main+xml"
_DANGLING_REL = (
    '<Relationship Id="rIdCustomUI" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customUI" '
    'Target="userCustomization/customUI.xml"/>'
)
_EXTERNAL_REL = (
    '<Relationship Id="rId99" TargetMode="External" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
    'Target="javascript:void(0);"/>'
)


def _build_docx(
    path: Path, *, paragraphs: list[str] | None = None,
    table_rows: list[list[str]] | None = None, tail: list[str] | None = None,
) -> bytes:
    """造一份正常的 docx：段落 → 表格 → 段落。"""
    document = Document()
    for text in paragraphs or []:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for col_index, value in enumerate(row):
                table.cell(row_index, col_index).text = value
    for text in tail or []:
        document.add_paragraph(text)
    document.save(path)
    return path.read_bytes()


def _patch_parts(data: bytes, patch: dict[str, str]) -> bytes:
    """按部件名替换 zip 内的文本部件。"""
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(buffer, "w") as target:
        for item in source.infolist():
            blob = source.read(item.filename)
            replacement = patch.get(item.filename)
            target.writestr(item.filename, blob if replacement is None else replacement)
    return buffer.getvalue()


class TestRepairDocxBytes:
    def test_healthy_package_is_returned_unchanged(self, tmp_path: Path) -> None:
        data = _build_docx(tmp_path / "ok.docx", paragraphs=["第一条　正文"])

        repaired, notes = repair_docx_bytes(data)

        assert notes == []
        assert repaired is data

    def test_dangling_internal_relationship_is_removed(self, tmp_path: Path) -> None:
        data = _build_docx(tmp_path / "dangling.docx", paragraphs=["第一条　正文"])
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            rels = archive.read("_rels/.rels").decode("utf-8")
        broken = _patch_parts(data, {
            "_rels/.rels": rels.replace("</Relationships>", _DANGLING_REL + "</Relationships>")
        })

        with pytest.raises(Exception):
            Document(io.BytesIO(broken))

        repaired, notes = repair_docx_bytes(broken)

        assert any("userCustomization/customUI.xml" in note for note in notes)
        assert Document(io.BytesIO(repaired)).paragraphs[0].text == "第一条　正文"

    def test_external_hyperlink_relationship_is_kept(self, tmp_path: Path) -> None:
        """法律库的 javascript: 外链不是缺失部件，不能被误删。"""
        data = _build_docx(tmp_path / "links.docx", paragraphs=["第一条　正文"])
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            rels = archive.read("_rels/.rels").decode("utf-8")
        with_links = _patch_parts(data, {
            "_rels/.rels": rels.replace("</Relationships>", _EXTERNAL_REL + "</Relationships>")
        })

        repaired, notes = repair_docx_bytes(with_links)

        assert notes == []
        assert repaired is with_links

    def test_macro_enabled_main_type_is_rewritten(self, tmp_path: Path) -> None:
        data = _build_docx(tmp_path / "macro.docx", paragraphs=["第一条　正文"])
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            content_types = archive.read("[Content_Types].xml").decode("utf-8")
        broken = _patch_parts(data, {
            "[Content_Types].xml": content_types.replace(_WML_MAIN, _MACRO_MAIN)
        })

        with pytest.raises(Exception):
            Document(io.BytesIO(broken))

        repaired, notes = repair_docx_bytes(broken)

        assert any("macro-enabled" in note for note in notes)
        assert Document(io.BytesIO(repaired)).paragraphs[0].text == "第一条　正文"


class TestLoaderTolerance:
    def test_loads_document_with_dangling_relationship(self, tmp_path: Path) -> None:
        data = _build_docx(tmp_path / "d.docx", paragraphs=["第一条　正文"])
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            rels = archive.read("_rels/.rels").decode("utf-8")
        path = tmp_path / "dangling.docx"
        path.write_bytes(_patch_parts(data, {
            "_rels/.rels": rels.replace("</Relationships>", _DANGLING_REL + "</Relationships>")
        }))

        result = DocxLoader().load(str(path))

        assert "第一条　正文" in result.content
        assert result.metadata["docx_repairs"]

    def test_loads_macro_enabled_document(self, tmp_path: Path) -> None:
        data = _build_docx(tmp_path / "d.docx", paragraphs=["第一条　正文"])
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            content_types = archive.read("[Content_Types].xml").decode("utf-8")
        path = tmp_path / "macro.docx"
        path.write_bytes(_patch_parts(data, {
            "[Content_Types].xml": content_types.replace(_WML_MAIN, _MACRO_MAIN)
        }))

        result = DocxLoader().load(str(path))

        assert "第一条　正文" in result.content
        assert result.metadata["docx_repairs"]

    def test_unrepairable_file_still_raises(self, tmp_path: Path) -> None:
        path = tmp_path / "junk.docx"
        path.write_bytes("这不是一个 zip".encode("utf-8"))

        with pytest.raises(ValueError):
            DocxLoader().load(str(path))

    def test_healthy_document_has_no_repair_metadata(self, tmp_path: Path) -> None:
        path = tmp_path / "ok.docx"
        _build_docx(path, paragraphs=["第一条　正文"])

        result = DocxLoader().load(str(path))

        assert "docx_repairs" not in result.metadata


class TestTableText:
    def test_table_text_is_kept_in_document_order(self, tmp_path: Path) -> None:
        path = tmp_path / "ordered.docx"
        _build_docx(
            path,
            paragraphs=["第一条　正文"],
            table_rows=[["第二条　表格里的条文"], ["第三条　仍在表格里"]],
            tail=["第四条　末尾"],
        )

        content = DocxLoader().load(str(path)).content

        assert content.split("\n\n") == [
            "第一条　正文",
            "第二条　表格里的条文",
            "第三条　仍在表格里",
            "第四条　末尾",
        ]

    def test_table_only_document_is_not_empty(self, tmp_path: Path) -> None:
        """实测 37 份文档整部法条都在一张表格里，此前取到 0 字符。"""
        path = tmp_path / "table_only.docx"
        _build_docx(path, table_rows=[["第一条　整部法条"], ["第二条　都在表格里"]])

        result = DocxLoader().load(str(path))

        assert result.content.strip()
        assert "第一条　整部法条" in result.content

    def test_merged_cells_are_not_duplicated(self, tmp_path: Path) -> None:
        document = Document()
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "合并后的文字"
        table.cell(0, 0).merge(table.cell(0, 1))
        path = tmp_path / "merged.docx"
        document.save(path)

        content = DocxLoader().load(str(path)).content

        assert content == "合并后的文字"

    def test_every_row_of_a_large_table_is_kept(self, tmp_path: Path) -> None:
        """去重曾用裸 ``id(tc)``，lxml 代理 id 复用导致整行文字被丢弃。

        实测语料里《吉林省个体工商户条例》有 125 行的单列表格，整部法条都在里面，
        修复前只留下第一行（687 字符 / 5288 字符）。
        """
        rows = [[f"第{n}条　本条内容编号{n}"] for n in range(1, 61)]
        path = tmp_path / "wide.docx"
        _build_docx(path, table_rows=rows)

        content = DocxLoader().load(str(path)).content

        for n in range(1, 61):
            assert f"第{n}条　本条内容编号{n}" in content
        assert len(content.split("\n\n")) == 60

    def test_nested_table_text_is_kept(self, tmp_path: Path) -> None:
        """实测 11 份文档是"表格套表格"，外层单元格本身没有文字。"""
        document = Document()
        outer = document.add_table(rows=1, cols=1)
        outer.cell(0, 0).paragraphs[0].text = "第一条　外层正文"
        inner = outer.cell(0, 0).add_table(rows=2, cols=1)
        inner.cell(0, 0).text = "第二条　内层表格条文"
        inner.cell(1, 0).text = "第三条　内层表格条文"
        path = tmp_path / "nested.docx"
        document.save(path)

        content = DocxLoader().load(str(path)).content

        assert content.split("\n\n") == [
            "第一条　外层正文",
            "第二条　内层表格条文",
            "第三条　内层表格条文",
        ]

    def test_text_box_body_falls_back_to_xml(self, tmp_path: Path) -> None:
        """文字在文本框里时 python-docx 结构遍历取不到，必须退回 XML 兜底。

        实测《清远市城市市容和环境卫生管理条例》整部法条（49 条 / 9,659 字）都在
        文本框里，此前 loader 返回 0 字符。
        """
        document = Document()
        document.add_paragraph("占位段落")
        path = tmp_path / "textbox.docx"
        document.save(path)
        # 把占位段落换成一个只含文本框的段落：python-docx 读不到 txbxContent 里的文字。
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
        start = xml.index("<w:p")
        end = xml.index("</w:p>") + len("</w:p>")
        text_box = (
            '<w:p xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
            "<w:r><mc:AlternateContent><mc:Fallback>"
            '<w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml">'
            "<v:textbox><w:txbxContent xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\">"
            "<w:p><w:r><w:t>第一条　文本框里的条文</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>第二条　仍在文本框里</w:t></w:r></w:p>"
            "</w:txbxContent></v:textbox></v:shape></w:pict>"
            "</mc:Fallback></mc:AlternateContent></w:r></w:p>"
        )
        path.write_bytes(_patch_parts(path.read_bytes(), {
            "word/document.xml": xml[:start] + text_box + xml[end:]
        }))

        content = DocxLoader().load(str(path)).content

        assert "第一条　文本框里的条文" in content
        assert "第二条　仍在文本框里" in content
